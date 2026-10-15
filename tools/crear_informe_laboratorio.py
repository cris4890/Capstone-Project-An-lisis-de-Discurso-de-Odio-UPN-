"""Construye el informe a partir de los resultados reales del laboratorio."""
from pathlib import Path
import html
import json
import shutil
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "entregables/laboratorio_unitarias"
EV = OUT / "evidencias"
summary = json.loads((EV / "resumen.json").read_text(encoding="utf-8"))
suite = ET.parse(EV / "suite_completa.xml").find("testsuite")
assert suite is not None
total, failures, errors = [int(suite.get(k, "0")) for k in ["tests", "failures", "errors"]]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Title"].font.size = Pt(25)
doc.styles["Heading 1"].font.size = Pt(16)
doc.styles["Heading 2"].font.size = Pt(12)
for element in doc.styles.element.xpath(".//w:pBdr"):
    element.getparent().remove(element)
header = section.header.paragraphs[0]
header.text = "UPN  |  HateDetect  |  Laboratorio de pruebas unitarias"
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.text = "14 de octubre de 2026                                                 Página "
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def p(text):
    return doc.add_paragraph(text)


def h(text):
    return doc.add_heading(text, level=1)


def code(text):
    paragraph = p(text)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def table(headers, rows):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Light Shading Accent 1"
    for cell, text in zip(tab.rows[0].cells, headers):
        cell.text = text
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8EDF2")
        cell._tc.get_or_add_tcPr().append(shd)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
    repeat = OxmlElement("w:tblHeader")
    tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = tab.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
        no_split = OxmlElement("w:cantSplit")
        cells[0]._tc.getparent().get_or_add_trPr().append(no_split)
    for row in tab.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(5)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    return tab


doc.add_heading("Pruebas unitarias de HateDetect", 0)
p("Informe de laboratorio adaptado al proyecto en Python")
p("Proyecto: Análisis y clasificación automática del discurso de odio en comunidades de Reddit r/Millennials y r/GenZ.")
p("Presentado por Jeanpierre Amilcar Casimiro Guerra. Universidad Privada del Norte. Fecha de ejecución local: 14 de octubre de 2026.")
h("Objetivo y resultado")
p("Verificamos diez comportamientos unitarios de HateDetect mediante pytest. La ejecución inicial aprobó los diez casos. Al introducir un defecto en una copia aislada, una prueba detectó la selección incorrecta del modelo. Después de restaurar esa copia, los diez casos volvieron a aprobarse. Los archivos originales conservaron sus hashes SHA256.")
p(f"La suite completa registró {total} pruebas, {failures} fallos y {errors} errores. Esta suite combina pruebas unitarias y otras comprobaciones de integración e interfaz; no se presenta su totalidad como pruebas unitarias.")
h("Adaptación de la guía del docente")
p("La guía proporcionada utiliza Java, Maven y JUnit en Apache NetBeans 28 con un sistema de ventas. Este informe aplica sus objetivos de diseño, ejecución, detección de fallos y regresión al sistema real desarrollado en Python. La aceptación de esta equivalencia tecnológica corresponde al docente; no consta su aprobación en este informe.")
table(["Guía de laboratorio", "Aplicación en HateDetect"], [
    ["Java y lógica de ventas", "Python y funciones de limpieza y evaluación"],
    ["JUnit y métodos @Test", "pytest y diez funciones test_p01 a test_p10"],
    ["assertEquals y assertThrows", "assert, pytest.approx y pytest.raises"],
    ["Maven Surefire y reportes XML", "pytest con --junitxml y registros TXT"],
    ["Defecto en el cálculo de IGV", "Mutación max a min en select_model"],
])
h("Alcance de las pruebas")
p("Los casos invocan funciones directamente, sin abrir Streamlit, conectar con Reddit ni acceder a una base de datos. Cubren resultados normales, entradas límite, excepciones y una métrica numérica. No miden carga concurrente ni la exactitud del clasificador sobre un corpus real.")

doc.add_page_break()
h("Matriz de diez casos unitarios")
p("Las entradas, resultados esperados y observados proceden de las propiedades del reporte XML de la ejecución final. La matriz completa también se entrega en CSV. Cadena vacía se representa como ''.")
cases = summary["ejecuciones"][-1]["casos"]
table(["Caso", "Entrada", "Esperado", "Observado", "Estado"], [
    [c["caso"], c["entrada"], c["esperado"], c["observado"], c["estado"]] for c in cases
])
h("Criterios de comprobación")
p("P01 a P05 verifican la limpieza y conservación del texto, incluido el ocultamiento de un correo ficticio. P06 exige seleccionar la puntuación de validación más alta. P07 y P08 comprueban que las entradas inválidas generen ValueError.")
p("P09 verifica FNR = FN / (TP + FN): un falso negativo entre dos ejemplos reales de Odio produce 0,5. Utiliza pytest.approx con tolerancia absoluta de 10⁻⁹. P10 comprueba que, sin ejemplos reales de Odio, su FNR sea None y no un cero que sugiera desempeño perfecto.")
p("P10 genera una advertencia de scikit-learn por trabajar con una sola clase. La ejecución registra esa advertencia y la prueba pasa; no se oculta ni se confunde con un error.")

doc.add_page_break()
h("Demostración del fallo y la corrección")
table(["Ejecución", "Aprobadas", "Fallidas", "Salida"], [[e["fase"], sum(c["estado"] == "PASÓ" for c in e["casos"]), sum(c["estado"] == "FALLÓ" for c in e["casos"]), e["codigo_salida"]] for e in summary["ejecuciones"]])
p("El ejecutor copia data_pipeline.py, experiment.py y los diez casos a una carpeta temporal. Solo allí sustituye la selección del máximo por la del mínimo. La mutación es un defecto didáctico, no un defecto que estuviera presente en producción.")
code("Correcto: return max(validation_scores, key=validation_scores.get)\nDefecto:  return min(validation_scores, key=validation_scores.get)")
p("Con modelo_a = 0,80 y modelo_b = 0,30, P06 esperaba modelo_a. La versión alterada devolvió modelo_b y pytest produjo AssertionError. El registro 02_fallo_controlado.txt conserva la traza completa.")
code("E  AssertionError: assert 'modelo_b' == 'modelo_a'\n1 failed, 9 passed")
p("La tercera ejecución restauró el contenido original de experiment.py en la copia temporal. P06 volvió a pasar y el resultado final fue de diez casos aprobados. El archivo resumen.json contiene los hashes de las fuentes y fuentes_intactas = true.")
h("Código de la prueba que detecta el defecto")
code('def test_p06_elige_mayor_validacion(record_property):\n    scores = {"modelo_a": 0.80, "modelo_b": 0.30}\n    actual = select_model(scores)\n    evidence(record_property, "P06", scores, "modelo_a", actual)\n    assert actual == "modelo_a"')
h("Interpretación de las aserciones")
p("assert compara el resultado observado con el esperado. pytest.approx permite comparar valores numéricos considerando tolerancia. pytest.raises verifica tanto el tipo de excepción como, en estos casos, parte de su mensaje. Si no ocurre la excepción prevista, la prueba falla.")
p("Una prueba unitaria revisa una función de manera acotada; una prueba de integración comprueba la colaboración entre componentes. Las pruebas de rendimiento estudian latencia, volumen o concurrencia. Los tiempos registrados por pytest no equivalen a una prueba de carga con JMeter.")

doc.add_page_break()
h("Reproducción y evidencias")
p("Desde la raíz del repositorio, con sus dependencias instaladas, se pueden reproducir los resultados mediante los siguientes comandos de PowerShell.")
code('.\\.venv\\Scripts\\python.exe -m pytest tests/test_laboratorio_unitarias.py -v -o junit_family=xunit1 --junitxml=resultado_laboratorio.xml')
code('.\\.venv\\Scripts\\python.exe tools/ejecutar_laboratorio.py')
code('.\\.venv\\Scripts\\python.exe -m pytest -q -o junit_family=xunit1 --junitxml=resultado_completo.xml')
p("El ejecutor sobrescribe únicamente los reportes de su carpeta de salida. Para conservar una ejecución anterior puede indicarse --output seguido de otra carpeta. La copia temporal se elimina al terminar y no contiene corpus, cuentas ni bases de datos de evaluadores.")
table(["Archivo o carpeta", "Contenido verificable"], [
    ["tests/test_laboratorio_unitarias.py", "Diez casos y propiedades de entrada y salida"],
    ["tools/ejecutar_laboratorio.py", "Ejecución aislada y validación de las tres fases"],
    ["evidencias/01_correcto y 03_restaurado", "Reportes TXT y XML con diez casos aprobados"],
    ["evidencias/02_fallo_controlado", "Reportes TXT y XML con el fallo P06"],
    ["evidencias/matriz_final.csv y resumen.json", "Matriz y trazabilidad mediante SHA256"],
    ["evidencias/suite_completa.xml y .txt", f"Suite completa de {total} pruebas"],
    ["evidencias.html", "Vista local del código y los registros originales"],
])
h("Correspondencia y límites del entregable")
p("Se entregan las fuentes, la matriz, los reportes y el ciclo de éxito, fallo y regresión. Las evidencias son registros reales de pytest. La vista HTML permite consultarlos junto con el código; no se presenta como una captura de NetBeans. Las capturas de pantalla del entorno solicitadas por la guía deben complementarse si el docente las exige, pues no hubo una superficie de navegador o escritorio disponible para capturarlas en esta sesión.")
p("Las actividades específicas de ventas, como descuentos o envío, no se incorporan a HateDetect porque no pertenecen a sus funciones. Tampoco se atribuye cumplimiento literal a la estructura Maven. Se solicita aceptar la adaptación a la tecnología del proyecto.")
p("Los resultados comprueban únicamente los comportamientos cubiertos. No resuelven los falsos negativos de discurso de odio detectados previamente ni acreditan la calidad del modelo con datos reales. Esa corrección y su evaluación mantienen un alcance separado.")
h("Fuente y trazabilidad")
p("Fuente académica: Guía práctica de pruebas unitarias con Apache NetBeans 28, documento proporcionado por el docente. Fuente experimental: ejecuciones locales guardadas en la carpeta evidencias. Python utilizado: " + summary["python"].split(" (")[0] + ".")
doc.save(OUT / "Informe_pruebas_unitarias_HateDetect.docx")

parts = ['<!doctype html><html lang="es"><meta charset="utf-8"><title>Evidencias del laboratorio HateDetect</title><style>body{font:16px system-ui;max-width:1100px;margin:40px auto;padding:20px;color:#17202a}pre{background:#f1f3f5;padding:20px;overflow:auto;font:13px Consolas,monospace}h2{margin-top:40px}</style><h1>Evidencias del laboratorio HateDetect</h1><p>Fuentes y registros reales de ejecución. Esta vista no es una captura de pantalla de un IDE.</p>']
sources = ["data_pipeline.py", "experiment.py", "tests/test_laboratorio_unitarias.py", "tools/ejecutar_laboratorio.py"]
parts.append("<h2>Estructura del laboratorio</h2><pre>" + html.escape("\n".join(sources)) + "</pre>")
for name in sources:
    parts.append("<h2>" + html.escape(name) + "</h2><pre>" + html.escape((ROOT / name).read_text(encoding="utf-8")) + "</pre>")
for name in ["01_correcto", "02_fallo_controlado", "03_restaurado"]:
    parts.append("<h2>" + name + "</h2><pre>" + html.escape((EV / (name + ".txt")).read_text(encoding="utf-8")) + "</pre>")
(OUT / "evidencias.html").write_text("\n".join(parts) + "</html>", encoding="utf-8")
for name in sources + ["requirements.lock.txt", "pytest.ini"]:
    target = OUT / "fuentes" / name
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / name, target)
(OUT / "LEEME.txt").write_text("Laboratorio adaptado a Python y pytest.\nAbra el informe Word y evidencias.html. Los registros TXT/XML y la matriz CSV están en evidencias.\nLas fuentes se incluyen como referencia; para reproducir use el repositorio completo y su entorno .venv.\nEjecute: .\\.venv\\Scripts\\python.exe tools/ejecutar_laboratorio.py\nEl fallo didáctico afecta únicamente una copia temporal.\nLa aceptación de la adaptación tecnológica y las capturas adicionales corresponde al docente.\n", encoding="utf-8")
print(OUT / "Informe_pruebas_unitarias_HateDetect.docx")
