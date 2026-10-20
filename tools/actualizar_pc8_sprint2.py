"""Actualiza una copia del informe PC8 con evidencias, sin alterar el original."""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument("source", type=Path)
parser.add_argument("output", type=Path)
args = parser.parse_args()
d = Document(args.source)
if "Title" not in d.styles:
    d.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH).base_style = d.styles["Normal"]
original = list(d.paragraphs)


def replace(index, text):
    original[index].text = text


replace(516, "El avance se describe por componentes y evidencias verificables. No se asigna un porcentaje global sin un denominador trazable de requisitos. El corpus autorizado, la anotación real, el entrenamiento definitivo y la evaluación científica final permanecen pendientes; la funcionalidad disponible no equivale al cierre de investigación.")
replace(521, "La ejecución del 19 de octubre de 2026 registró 105 pruebas aprobadas: 77 unitarias, 21 de integración y 7 de sistema mediante Streamlit AppTest. El inventario por caso y el JUnit XML se conservan en el paquete Sprint 2. Las pruebas técnicas no constituyen aceptación formal ni evaluación científica definitiva.")
replace(523, "Las pruebas unitarias verifican funciones acotadas de limpieza, anonimización, reglas, cálculos y validación con entradas controladas. Los recorridos de almacenamiento, respaldo, restauración y entrenamiento con persistencia se clasifican como integración. La clasificación se registra por caso mediante pytest y no se deduce únicamente del nombre del archivo.")
replace(529, "Los criterios de aceptación cubren clasificación, guardado voluntario, consulta de Historial, Reportes y visibilidad de errores y límites. Se dispone de verificación técnica y capturas reales. La Review de Sprint 2 está preparada: fecha de sesión, participantes efectivos, criterios aceptados y observaciones requieren registro humano. No se afirma que haya sido celebrada ni aprobada.")
replace(532, "GitHub Actions ejecutó satisfactoriamente las 105 pruebas de la versión 2812f23 de Sprint 2. Esta evidencia automatizada de calidad no constituye un despliegue productivo ni aceptación formal.")
replace(539, "La ficha KPI-PILOTO-001 del 04/10/2026 contiene 30 textos sintéticos y una clave de trabajo. El registro manual suma 203,84 segundos, 28 aciertos y 2 errores. La repetición de reglas v2 del 19/10/2026 mantiene 18 aciertos y 12 errores. Se midieron además treinta ventanas desde el clic hasta el resultado renderizado en el navegador: suma de 8,407 segundos y media de 0,280 segundos. La Tabla 24 distingue los alcances.")
replace(541, "Nota. Los tiempos E2E se toman desde el clic hasta un marcador nuevo de resultado y dos frames como aproximación al renderizado. Excluyen ingreso, revisión humana, guardado e intervalos entre casos. No son directamente equivalentes al trabajo manual. La capacidad y cobertura de ventanas se muestran de forma descriptiva; no acreditan ahorro laboral ni superioridad operativa causal.")
replace(545, "El piloto distingue velocidad de respuesta y calidad predictiva. El 60 % de exactitud de HateDetect es inferior al 93,33 % manual. La repetición E2E mejora el alcance del cronometraje frente a medir solo inferencia, pero no incluye todo el trabajo humano. No se calculan ahorro salarial ni retorno de inversión sin costos y ventanas comparables.")
replace(553, "Se dispone de una ficha piloto y de una repetición E2E para explorar los cinco KPI. Estas observaciones permiten detectar defectos del instrumento y del producto, pero no validan definitivamente el procedimiento. Faltan límites temporales equivalentes, repetición con evaluadores y muestra independiente.")
replace(556, "X. REFERENCIAS")
for p in original:
    if "Los tiempos de HateDetect corresponden al tiempo interno" in p.text:
        p.text = original[541].text
for idx in [8]:
    original[idx].style = d.styles["Title"]

# Actualizar las dos tablas que repetían las cifras internas no comparables.
rows = [
    ["Tiempo medio", "6,795 s/texto", "0,280 s E2E/texto", "Ventanas distintas; no calcular ahorro causal"],
    ["Capacidad sobre tiempos registrados", "8,83 textos/min", "214,11 textos/min de ventanas E2E", "Excluye ingreso y revisión; no es productividad integral"],
    ["Exactitud", "93,33 % (28/30)", "60,00 % (18/30)", "HateDetect: 33,33 puntos porcentuales menos"],
    ["Tasa de error", "6,67 % (2/30)", "40,00 % (12/30)", "HateDetect: 33,33 puntos porcentuales más"],
    ["Cobertura acumulada en 60 s", "26,67 % (8/30)", "100 % de ventanas (30/30)", "No es una sesión continua comparable; no equivale a recall"],
]
for index in [8, 23]:
    for target, values in zip(d.tables[index].rows[1:], rows):
        for cell, text in zip(target.cells, values):
            cell.text = text
d.tables[21].rows[8].cells[2].text = "105 pruebas aprobadas; inventario por tipo y GitHub Actions"
d.tables[21].rows[8].cells[3].text = "Aceptación formal y evaluación científica final"
d.tables[22].rows[2].cells[2].text = "Evidencias actualizadas: piloto, E2E, 105 pruebas, matriz, incidencias y rama Sprint 2"
d.tables[22].rows[2].cells[3].text = "Corpus autorizado y Review formal; no se declara cerrado"
d.tables[19].rows[3].cells[2].text = "Procedimiento documentado; cuenta, ACL, cifrado y actualizaciones no auditados"
for row in d.tables[25].rows[1:]:
    row.cells[2].text = "Verificación técnica; aceptación formal pendiente"
new = d.tables[14].add_row().cells
for cell, text in zip(new, ["R09", "Calidad predictiva insuficiente", "Piloto: 12/30 errores y 0 % recall Odio", "Alto; materializado en piloto"]):
    cell.text = text
new = d.tables[15].add_row().cells
for cell, text in zip(new, ["R09", "Mitigar y mantener revisión humana", "Preservar fallos, corregir con regresiones y evaluar corpus reservado", "Equipo; ratificación en Review"]):
    cell.text = text

# Insertar nuevos contenidos antes de Referencias respetando estilos del informe.
anchor = original[556]
def add(text, style="Normal"):
    p = anchor.insert_paragraph_before(text)
    p.style = d.styles[style]
    return p

add("7.7. Protocolo, auditoría y limitaciones", "Heading 2")
add("La hoja Prueba HateDetect contiene doce discrepancias entre D y J. El resumen original utiliza D, que coincide con la clave; J coincide con las salidas históricas. El anexo conserva ambos valores, normaliza Neutral a Neutro y usa las predicciones reproducidas. El original del Excel permanece sin modificaciones y se registra su SHA256.")
add("La muestra es pequeña, sintética y conocida; una repetición por texto no permite inferencia definitiva. No se acredita cegamiento del evaluador ni independencia de la clave. El navegador se encontraba cargado, se preservó el primer caso y se excluyeron los intervalos de inspección del cálculo de ventanas. El corpus real, su anotación y el Transformer definitivo siguen pendientes.")
add("Figura 5\nMatriz de confusión del piloto de treinta casos", "FIGURA")
p = add("")
p.add_run().add_picture(str(ROOT / "docs/sprint2/evidencias/matriz_confusion.png"), width=Inches(5.2))
add("Filas: referencia de trabajo. Columnas: predicción. Fuente: ejecución de reglas v2 sobre los 30 textos sintéticos del anexo; no representa desempeño poblacional.")
content = (ROOT / "docs/sprint2/capitulos.md").read_text(encoding="utf-8")
for title, start, end in [
    ("VIII. CONCLUSIONES Y TRABAJO PENDIENTE", "## VIII", "## IX"),
    ("IX. COMPROMISO ÉTICO", "## IX", "## X"),
]:
    h = add(title, "Heading 1")
    h.paragraph_format.page_break_before = True
    body = content.split(start, 1)[1].split(end, 1)[0].split("\n", 1)[1]
    for para in body.strip().split("\n\n"):
        p = add(para)
        p.paragraph_format.line_spacing = 1.5
anchor.paragraph_format.page_break_before = True

d.add_heading("Anexo E. Evidencias de Sprint 2", level=2)
d.add_paragraph("Paquete asociado: docs/sprint2 del repositorio, rama codex/sprint2-evidencias. Incluye anexo_kpi.csv con treinta textos, referencia, clasificación manual, salidas de las columnas D y J, nueva clasificación y tiempos; e2e_navegador.json; matriz_confusion.csv; errores.csv; pruebas.xml; inventario_pruebas.csv; version.json; trazabilidad, riesgos y protocolo. Los registros manuales se conservan como fueron aportados.")
d.add_paragraph("La Review está preparada y pendiente de celebración. Debe registrar fecha real, participantes efectivos, versión demostrada, criterios aceptados y observaciones. Los resultados automatizados no sustituyen esas decisiones.")
for title, filename in [("Figura 6\nResultado real del caso T003", "resultado_T003.png"), ("Figura 7\nHistorial aislado de demostración", "historial.png"), ("Figura 8\nReportes del corpus sintético", "reportes_grafico.png")]:
    d.add_paragraph(title, style="FIGURA")
    d.add_picture(str(ROOT / "docs/sprint2/capturas" / filename), width=Inches(5.7))
    d.add_paragraph("Captura real del navegador, 19/10/2026. Instancia local aislada; no acredita aceptación formal ni datos reales.")

# Corregir numeración de tablas por orden de aparición, incluidas referencias en prosa.
captions = [p for p in d.paragraphs if p.style.name == "TABLA" and re.match(r"Tabla \d+", p.text)]
mapping = {re.match(r"Tabla (\d+)", p.text).group(1): str(i) for i, p in enumerate(captions, 1)}
caption_ids = {id(p._p) for p in captions}
for p in d.paragraphs:
    if p.style.name != "TABLA" and "Tabla " in p.text and p.style.name != "table of figures":
        p.text = re.sub(r"Tabla (\d+)\b", lambda m: "Tabla " + mapping.get(m.group(1), m.group(1)), p.text)
for i, p in enumerate(captions, 1):
    p.text = re.sub(r"^Tabla \d+", f"Tabla {i}", p.text)

# Reemplazar índices antiguos con campos actualizables de Word.
for p in original[29:114]:
    if p._p.getparent() is not None:
        p._p.getparent().remove(p._p)
for child in list(d.element.body):
    if child.tag == qn("w:sdt") and " TOC " in etree.tostring(child, encoding="unicode"):
        d.element.body.remove(child)
index_anchor = original[114]
original[28].text = "ÍNDICE GENERAL"
def field_before(instruction):
    p = index_anchor.insert_paragraph_before("")
    r = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); r.append(begin); p._p.append(r)
    r = OxmlElement("w:r"); code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = instruction; r.append(code); p._p.append(r)
    for kind in ["separate", "end"]:
        r = OxmlElement("w:r"); element = OxmlElement("w:fldChar"); element.set(qn("w:fldCharType"), kind); r.append(element); p._p.append(r)
    return p
field_before(' TOC \\o "1-3" \\h \\z \\u ')
p = index_anchor.insert_paragraph_before("ÍNDICE DE TABLAS"); p.style = "Heading 1"; p.paragraph_format.page_break_before = True
field_before(' TOC \\h \\z \\t "TABLA;1" ')
p = index_anchor.insert_paragraph_before("ÍNDICE DE FIGURAS"); p.style = "Heading 1"; p.paragraph_format.page_break_before = True
field_before(' TOC \\h \\z \\t "FIGURA;1" ')
index_anchor.paragraph_format.page_break_before = True
for p in d.paragraphs:
    if p.style.name in {"TABLA", "FIGURA"}:
        if not re.match(r"^(Tabla|Figura) \d+", p.text):
            p.style = d.styles["Normal"]
            continue
        p.paragraph_format.keep_with_next = True
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.style.name in {"TABLA", "FIGURA"}:
                    p.style = d.styles["Normal"]
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.paragraph_format.keep_with_next = True
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
args.output.parent.mkdir(parents=True, exist_ok=True)
d.save(args.output)
print(args.output)
